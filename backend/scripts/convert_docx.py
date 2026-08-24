"""将 Word 题库转换为平台导入模板 xlsx。

用法：
    uv run python scripts/convert_docx.py <input.docx> [output.xlsx]

解析约定：
- 单选/多选：以 "数字." 开头识别题干，"A."~"D." 行为选项，"答案：" 行为答案。
- 判断：题干同上；答案行为 √/×（或 对/错、正/误），转换为 正确/错误。

源文档中缺失的 解析/难度/标签/分值/分组ID 列分别填空/2/空/2/空。
表头复用 app.utils.excel.HEADERS，确保与下载的导入模板完全一致。
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.utils.excel import HEADERS, SHEET_ORDER

HEADER_FILL = PatternFill("solid", fgColor="E60012")
HEADER_FONT = Font(color="FFFFFF", bold=True)

Q_RE = re.compile(r"^\d+[.、．]\s*(.+)$")
OPT_RE = re.compile(r"^([A-Za-z])[.、:：]\s*(.+)$")
ANS_RE = re.compile(r"^答案\s*[：:]\s*(.+)$")

# 判断题符号 → 正确/错误
TRUE_TOKENS = {"√", "对", "正确", "正", "T", "true", "是", "A"}
FALSE_TOKENS = {"×", "x", "X", "错", "错误", "误", "F", "false", "否", "B"}


def parse_docx(path: Path) -> dict[str, list[dict]]:
    """按题型分组解析，返回 {sheet_name: [question, ...]}。"""
    import docx

    doc = docx.Document(str(path))
    lines = [p.text.strip() for p in doc.paragraphs]

    sections: dict[str, list[dict]] = {s: [] for s in SHEET_ORDER}
    cur_sheet: str | None = None
    cur_q: dict | None = None

    def flush() -> None:
        nonlocal cur_q
        if cur_q and cur_sheet:
            if cur_q.get("question") and cur_q.get("answer_raw") is not None:
                sections[cur_sheet].append(cur_q)
        cur_q = None

    for line in lines:
        if not line:
            continue
        # 章节标题：一、单选题 / 二、多选题 / 三、判断题
        if re.match(r"^[一二三四五六七八九十]、", line):
            flush()
            for st in SHEET_ORDER:
                if st in line:
                    cur_sheet = st
                    break
            continue
        if cur_sheet is None:
            continue

        if ANS_RE.match(line):
            if cur_q is not None:
                cur_q["answer_raw"] = ANS_RE.match(line).group(1).strip()
            continue

        if Q_RE.match(line):
            flush()
            cur_q = {"question": Q_RE.match(line).group(1).strip(), "options": [],
                     "answer_raw": None}
            continue

        if cur_q is not None:
            m = OPT_RE.match(line)
            if m:
                cur_q["options"].append(m.group(2).strip())

    flush()
    return sections


def _normalize_options(options: list[str]) -> str:
    """选项重新按 A./B./C./D. 顺序编号，换行连接。"""
    letters = "ABCDEFGHIJKLMNOP"
    parts = [f"{letters[i]}.{o}" for i, o in enumerate(options)]
    return "\n".join(parts)


def _judge_answer(raw: str) -> str | None:
    r = raw.strip()
    if r in TRUE_TOKENS:
        return "正确"
    if r in FALSE_TOKENS:
        return "错误"
    # 兼容 "正确"/"错误" 文本
    if "正" in r and "错" not in r:
        return "正确"
    if "错" in r or "误" in r:
        return "错误"
    return None


def build_xlsx(sections: dict[str, list[dict]], out: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    widths = [40, 50, 20, 30, 8, 20, 8, 12]

    for st in SHEET_ORDER:
        ws = wb.create_sheet(st)
        headers = HEADERS[st]
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for i, w in enumerate(widths[: len(headers)], 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        for q in sections.get(st, []):
            row = _build_row(st, q)
            ws.append(row)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))


def _build_row(sheet_type: str, q: dict) -> list:
    """按表头顺序构造一行；缺失列填默认值。"""
    headers = HEADERS[sheet_type]
    row: list = [""] * len(headers)

    def set_col(name: str, value) -> None:
        if name in headers:
            row[headers.index(name)] = value

    set_col("题干", q["question"])
    set_col("解析", "")
    set_col("难度", 2)
    set_col("知识点标签", "")
    set_col("分值", 2)
    set_col("所属分组ID", "")

    ans_raw = q.get("answer_raw") or ""
    if sheet_type in ("单选题", "多选题"):
        set_col("选项", _normalize_options(q.get("options", [])))
        set_col("答案", ans_raw.upper().replace(" ", ""))
    elif sheet_type == "判断题":
        set_col("答案", _judge_answer(ans_raw) or "")
    elif sheet_type == "简答题":
        set_col("参考答案", ans_raw)
    elif sheet_type == "拖拽题":
        set_col("题项与正确容器", ans_raw)
    return row


def main() -> None:
    if len(sys.argv) < 2:
        print("用法: python scripts/convert_docx.py <input.docx> [output.xlsx]")
        sys.exit(1)
    src = Path(sys.argv[1]).resolve()
    if not src.exists():
        print(f"文件不存在: {src}")
        sys.exit(1)
    dst = Path(sys.argv[2]).resolve() if len(sys.argv) > 2 else src.with_suffix(".xlsx")

    sections = parse_docx(src)
    total = sum(len(v) for v in sections.values())
    for st in SHEET_ORDER:
        print(f"  {st}: {len(sections[st])} 题")
    print(f"合计: {total} 题")

    build_xlsx(sections, dst)
    print(f"已生成导入模板: {dst}")


if __name__ == "__main__":
    main()

"""本轮评审 6 个 critical 缺陷的回归测试。

每个用例都对应一个已修复的缺陷，断言的是**修复后的行为**而非实现细节：

1. 手工选题 / 模板试卷必须按题目自身分值固化（否则 total_score 与及格判定双错）；
2. 部门管理员用范围内题库即可创建规则组卷考试（不再强制要求 rules["group_ids"]）；
3. 登录对"邮箱不存在"也要付一次 bcrypt 开销（消除账号枚举的计时侧信道）；
4. 表头不一致的行不得进入导入队列（列按位置读取，错列会静默污染数据）；
5. 邮箱归一迁移在大小写变体重复时必须能完成、不丢数据、可重复执行；
6. 删除题目 / 题库后必须重算受影响的每日聚合（否则概览统计永久偏高）。
"""

from __future__ import annotations

import importlib.util
import secrets
import sqlite3
import sys
from io import BytesIO
from pathlib import Path

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.deps import dept_scope_ids
from app.core.errors import DomainError
from app.core.security import hash_password
from app.database import db_session, init_db
from app.models.exam import ExamDefinition, PaperTemplate
from app.models.group import Group
from app.models.question import Question, QuestionBank
from app.models.record import PracticeRecord
from app.models.stats import StatsUserDaily
from app.models.user import User
from app.schemas.exam import ExamCreateIn
from app.services import auth_service, exam_service, import_service, practice_service, question_service


# ---------- 公共构造 ----------
def _mk_user(db, role: str = "user", dept_group_id: int | None = None) -> User:
    user = User(
        email=f"{secrets.token_hex(4)}@example.com",
        password_hash=hash_password("pw123456"),
        name="用户",
        role=role,
        status="active",
        email_verified=True,
        dept_group_id=dept_group_id,
    )
    db.add(user)
    db.flush()
    return user


def _mk_bank(db, *, group_id: int | None = None, score: float = 2.0, answer: str = "A"):
    bank = QuestionBank(name="题库", group_id=group_id, practice_enabled=True)
    db.add(bank)
    db.flush()
    question = Question(
        bank_id=bank.id,
        type="单选题",
        question="1+1=?",
        options=["A.1", "B.2"],
        answer=answer,
        analysis="",
        difficulty=1,
        tags=[],
        score=score,
        group_id=group_id,
    )
    db.add(question)
    db.flush()
    return bank, question


def _mk_exam(db, *, creator_id: int, **overrides) -> ExamDefinition:
    fields = {
        "name": "考试",
        "type": "formal",
        "rules": {},
        "group_ids": None,
        "duration_min": 30,
        "pass_score": 60,
        "max_attempts": 0,
        "status": "published",
        "created_by": creator_id,
    }
    fields.update(overrides)
    exam = ExamDefinition(**fields)
    db.add(exam)
    db.flush()
    return exam


def _take_exam(db, user: User, exam_id: int, qid: int, answer: str = "A") -> dict:
    """开考 → 作答 → 交卷，返回 submit_exam 的结果。"""
    payload = exam_service.start_exam(db, user, exam_id)
    session_id = payload["session_id"]
    exam_service.submit_answer(db, user, session_id, qid, answer, 1)
    return {"session": payload, "result": exam_service.submit_exam(db, user, session_id)}


# ---------- 1. 分值固化 ----------
def test_manual_exam_freezes_question_own_score():
    """手工选题考试必须按题目分值固化，而非一律 2.0（回归：判分错误）。"""
    init_db()
    with db_session() as db:
        student = _mk_user(db)
        _, question = _mk_bank(db, score=5.0)
        exam = _mk_exam(db, creator_id=student.id, name="手工卷", manual_questions=[question.id])
        db.commit()
        student_id, exam_id, qid = student.id, exam.id, question.id

    with db_session() as db:
        student = db.get(User, student_id)
        out = _take_exam(db, student, exam_id, qid)

    assert out["session"]["questions"][0]["score"] == 5.0, "固化分应取题目自身的 5.0"
    assert out["result"]["total_score"] == 5.0, "卷面总分应为 5.0（原实现恒为 2.0）"
    assert out["result"]["score"] == 5.0
    assert out["result"]["passed"] is True


def test_template_exam_freezes_question_own_score():
    """模板已含题目清单的路径同样必须按题目分值固化。"""
    init_db()
    with db_session() as db:
        admin = _mk_user(db, role="super_admin")
        _, question = _mk_bank(db, score=7.5)
        template = PaperTemplate(
            name="模板",
            mode="formal",
            config={"type_quota": {"单选题": 1}},
            question_ids=[question.id],
            created_by=admin.id,
        )
        db.add(template)
        db.flush()
        exam = _mk_exam(db, creator_id=admin.id, name="模板卷", paper_template_id=template.id)
        db.commit()
        admin_id, exam_id, qid = admin.id, exam.id, question.id

    with db_session() as db:
        admin = db.get(User, admin_id)
        out = _take_exam(db, admin, exam_id, qid)

    assert out["session"]["questions"][0]["score"] == 7.5
    assert out["result"]["total_score"] == 7.5


# ---------- 2. 部门管理员组卷范围 ----------
def _mk_dept_with_admin(db, name: str = "研发部"):
    dept = Group(name=name, type="部门", parent_id=None)
    db.add(dept)
    db.flush()
    admin = _mk_user(db, role="dept_admin", dept_group_id=dept.id)
    return dept, admin


def test_dept_admin_can_create_rule_exam_with_in_scope_bank():
    """范围内题库即构成范围证据（回归：真实管理端请求恒 403）。"""
    init_db()
    with db_session() as db:
        dept, admin = _mk_dept_with_admin(db)
        bank, _ = _mk_bank(db, group_id=dept.id)
        scope = dept_scope_ids(db, admin)
        assert dept.id in scope

        payload = ExamCreateIn(
            name="规则卷",
            type="formal",
            group_ids=[dept.id],
            # 管理端真实提交的形状：只有 bank_ids，没有 rules["group_ids"]
            rules={"type_quota": {"单选题": 1}, "bank_ids": [bank.id], "order_mode": "bank"},
        )
        res = exam_service.create_exam(db, payload, admin, scope)

    assert res["id"], "部门管理员必须能创建规则组卷考试"


def test_dept_admin_rejected_when_bank_out_of_scope():
    """范围外题库必须仍然被拒（修复不能放宽为无校验）。"""
    init_db()
    with db_session() as db:
        _, admin = _mk_dept_with_admin(db, "研发部")
        other_dept = Group(name="市场部", type="部门", parent_id=None)
        db.add(other_dept)
        db.flush()
        foreign_bank, _ = _mk_bank(db, group_id=other_dept.id)
        scope = dept_scope_ids(db, admin)

        payload = ExamCreateIn(
            name="越权卷",
            type="formal",
            group_ids=[admin.dept_group_id],
            rules={"type_quota": {"单选题": 1}, "bank_ids": [foreign_bank.id]},
        )
        with pytest.raises((DomainError, HTTPException)) as exc:
            exam_service.create_exam(db, payload, admin, scope)

    assert exc.value.status_code == 403


def test_dept_admin_rejected_when_no_source_scope_evidence():
    """既无 bank_ids 也无 rules["group_ids"] 时无法证明范围，fail-closed 拒绝。"""
    init_db()
    with db_session() as db:
        dept, admin = _mk_dept_with_admin(db)
        scope = dept_scope_ids(db, admin)
        payload = ExamCreateIn(
            name="无来源卷",
            type="formal",
            group_ids=[dept.id],
            rules={"type_quota": {"单选题": 1}},
        )
        with pytest.raises((DomainError, HTTPException)) as exc:
            exam_service.create_exam(db, payload, admin, scope)

    assert exc.value.status_code == 403


def test_dept_admin_can_rename_exam_with_inherited_template():
    """存量配置未变时不得复检（回归：连改名都被 403 永久锁死）。"""
    init_db()
    with db_session() as db:
        dept, admin = _mk_dept_with_admin(db)
        super_admin = _mk_user(db, role="super_admin")
        template = PaperTemplate(
            name="全局模板", mode="formal", config={}, question_ids=None, created_by=super_admin.id
        )
        db.add(template)
        db.flush()
        exam = _mk_exam(db, creator_id=super_admin.id, name="旧名", paper_template_id=template.id, group_ids=[dept.id])
        db.commit()
        exam_id = exam.id
        scope = dept_scope_ids(db, admin)

        res = exam_service.update_exam(db, exam_id, {"name": "新名"}, scope)

    assert res["name"] == "新名"


# ---------- 3. 登录计时对齐 ----------
def test_login_pays_bcrypt_cost_for_unknown_email(monkeypatch):
    """未知邮箱分支必须同样执行一次密码校验（等价工作量），且响应不可区分。"""
    init_db()
    calls: list[str] = []
    real_verify = auth_service.verify_password

    def spy(plain: str, hashed: str) -> bool:
        calls.append(hashed)
        return real_verify(plain, hashed)

    monkeypatch.setattr(auth_service, "verify_password", spy)

    with db_session() as db:
        user = _mk_user(db)
        db.commit()
        known_email = user.email

    with db_session() as db:
        with pytest.raises((DomainError, HTTPException)) as unknown:
            auth_service.login(db, f"nobody-{secrets.token_hex(4)}@example.com", "whatever")
        unknown_calls = len(calls)
        calls.clear()
        with pytest.raises((DomainError, HTTPException)) as known:
            auth_service.login(db, known_email, "wrong-password")
        known_calls = len(calls)

    assert unknown.value.status_code == 401
    assert known.value.status_code == 401
    assert unknown_calls == 1, "未知邮箱必须付出与命中分支相同的 bcrypt 开销"
    assert known_calls == 1
    assert unknown.value.detail == known.value.detail, "两条分支的响应必须完全一致"


# ---------- 4. 导入表头校验 ----------
def _single_sheet_workbook(*, rename_header: bool) -> bytes:
    """只保留「单选题」Sheet 的模板副本（数据行取自真实模板），可选改其表头首格。"""
    from app.utils.excel import build_template

    wb = load_workbook(BytesIO(build_template().getvalue()))
    for name in list(wb.sheetnames):
        if name != "单选题":
            del wb[name]
    if rename_header:
        wb["单选题"].cell(row=1, column=1).value = "题干（被改名）"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def test_mismatched_header_rows_are_not_queued_for_import():
    """表头不一致的 Sheet 的行不得进入 valid_rows（回归：错列静默落库）。"""
    from app.utils.excel import build_template, parse_workbook

    mismatched = _single_sheet_workbook(rename_header=True)
    preview_obj = parse_workbook(BytesIO(mismatched))
    assert preview_obj.total == 0, "前提：parse_workbook 应拒绝表头不一致的 Sheet"
    assert preview_obj.errors

    # preview() 的待导入清单必须与 parse_workbook 的判定同源（现在共用 all_rows）
    res = import_service.preview(None, mismatched, None, None, "新库", 1, None)
    assert res["valid_count"] == 0, "表头不一致的行绝不能进入导入队列"

    # 对照：同一 Sheet 不改表头仍可正常导入（防止修复把正常文件一起拒掉）
    healthy = _single_sheet_workbook(rename_header=False)
    ok = import_service.preview(None, healthy, None, None, "新库", 1, None)
    assert ok["valid_count"] > 0
    # 完整模板（6 个 Sheet）也必须照常可用
    full = import_service.preview(None, build_template().getvalue(), None, None, "新库", 1, None)
    assert full["valid_count"] >= ok["valid_count"]


# ---------- 5. 邮箱归一迁移 ----------
_MIGRATION = Path(__file__).resolve().parent.parent / "scripts" / "migrate_2026_09_16.py"


def _load_migration():
    spec = importlib.util.spec_from_file_location("migrate_2026_09_16", _MIGRATION)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules["migrate_2026_09_16"] = module
    spec.loader.exec_module(module)
    return module


def _users_db(path: Path) -> sqlite3.Connection:
    conn = sqlite3.connect(str(path))
    conn.execute("CREATE TABLE users (id INTEGER PRIMARY KEY, email VARCHAR NOT NULL)")
    conn.execute("CREATE UNIQUE INDEX ix_users_email ON users (email)")
    return conn


def test_email_normalization_survives_case_variant_duplicates(tmp_path: Path):
    """大小写变体重复：必须完成归一、保留两行、且最终值互不冲突。"""
    migration = _load_migration()
    conn = _users_db(tmp_path / "dup.db")
    conn.execute("INSERT INTO users (id, email) VALUES (1, 'Admin@x.com')")
    conn.execute("INSERT INTO users (id, email) VALUES (2, 'ADMIN@X.com')")
    conn.commit()

    changed = migration.migrate_normalize_emails(conn, dry_run=False)

    rows = dict(conn.execute("SELECT id, email FROM users").fetchall())
    assert changed == 2
    assert len(rows) == 2, "归一不得丢行"
    assert rows[1] == "admin@x.com"
    assert rows[2] != rows[1], "冲突行必须被改名，否则撞唯一索引"
    assert rows[2].startswith("admin@x.com")
    assert len(set(rows.values())) == 2

    # 幂等：再跑一次应是 no-op
    assert migration.migrate_normalize_emails(conn, dry_run=False) == 0
    assert dict(conn.execute("SELECT id, email FROM users").fetchall()) == rows
    conn.close()


def test_email_normalization_noop_when_already_normalized(tmp_path: Path):
    migration = _load_migration()
    conn = _users_db(tmp_path / "clean.db")
    conn.execute("INSERT INTO users (id, email) VALUES (1, 'a@x.com')")
    conn.commit()

    assert migration.migrate_normalize_emails(conn, dry_run=False) == 0
    assert conn.execute("SELECT COUNT(*) FROM users").fetchone()[0] == 1
    conn.close()


# ---------- 6. 删除后的统计重算 ----------
def _total_answer_count(db, user_id: int) -> int:
    rows = db.execute(select(StatsUserDaily).where(StatsUserDaily.user_id == user_id)).scalars().all()
    return sum(row.answer_count for row in rows)


def test_delete_question_refreshes_daily_stats():
    """删除题目后聚合必须重算（回归：概览统计永久高于真实值）。"""
    init_db()
    with db_session() as db:
        user = _mk_user(db)
        _, question = _mk_bank(db)
        db.commit()
        user_id, qid = user.id, question.id

    with db_session() as db:
        practice_service.answer_question(db, user_id, qid, "A", "sequence")

    with db_session() as db:
        assert _total_answer_count(db, user_id) == 1, "前提：作答后聚合应有 1 条"

    with db_session() as db:
        question_service.delete_question(db, qid, scope=None)

    with db_session() as db:
        assert db.execute(select(PracticeRecord).where(PracticeRecord.question_id == qid)).first() is None
        assert _total_answer_count(db, user_id) == 0, "删除题目后聚合必须同步归零"


def test_delete_bank_refreshes_daily_stats():
    """删除题库（连带删题）同样必须重算聚合。"""
    init_db()
    with db_session() as db:
        user = _mk_user(db)
        bank, question = _mk_bank(db)
        db.commit()
        user_id, bank_id, qid = user.id, bank.id, question.id

    with db_session() as db:
        practice_service.answer_question(db, user_id, qid, "A", "sequence")
    with db_session() as db:
        assert _total_answer_count(db, user_id) == 1

    with db_session() as db:
        question_service.delete_bank(db, bank_id, scope=None)

    with db_session() as db:
        assert _total_answer_count(db, user_id) == 0, "删除题库后聚合必须同步归零"

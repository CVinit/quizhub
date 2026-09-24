"""Excel 解析单元测试。"""

from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.utils.excel import EXAMPLE_PREFIX, build_template, parse_workbook


def test_build_template_returns_xlsx():
    buf = build_template()
    data = buf.getvalue()
    assert data[:2] == b"PK"  # xlsx 是 zip


def test_parse_empty_workbook():
    wb = Workbook()
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    result = parse_workbook(buf)
    assert result.total == 0
    assert result.rows == []


def test_parse_template_skips_example_rows():
    """模板内置示例题（供用户参考格式）必须被跳过：否则用户不删示例行就会把示例题导入题库。"""
    buf = build_template()
    result = parse_workbook(buf)
    assert result.total == 0, "示例行不应进入导入队列"
    assert result.type_dist == {}
    assert result.errors == []

    # 示例行仍留在模板里（可读性不受影响），只是带上了可识别前缀
    wb = load_workbook(BytesIO(build_template().getvalue()))
    assert str(wb["单选题"].cell(row=2, column=1).value).startswith(EXAMPLE_PREFIX)


def test_parse_single_choice_row():
    wb = Workbook()
    ws = wb.active
    ws.title = "单选题"
    ws.append(["题干", "选项", "答案", "解析", "难度", "知识点标签", "分值", "所属分组ID"])
    ws.append(["HTTP 默认端口？", "21\n80\n443\n8080", "B", "HTTP默认80", 2, "网络", 2, ""])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    result = parse_workbook(buf)
    assert result.total == 1
    row = result.rows[0]
    assert row.type == "单选题"
    assert row.valid is True


def test_parse_single_choice_rejects_multi_letter_answer():
    """单选题答案必须是单个字母（与手动创建路径同口径），否则会存下永远判错的题。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "单选题"
    ws.append(["题干", "选项", "答案", "解析", "难度", "知识点标签", "分值", "所属分组ID"])
    ws.append(["HTTP 默认端口？", "21\n80\n443\n8080", "AB", "", 2, "网络", 2, ""])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    result = parse_workbook(buf)
    assert result.total == 1
    assert result.rows[0].valid is False
    assert "单选题答案必须是一个选项字母" in result.rows[0].error


def test_parse_multi_choice_accepts_multi_letter_answer():
    """多选题仍应接受多字母答案（且排序归一）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "多选题"
    ws.append(["题干", "选项", "答案", "解析", "难度", "知识点标签", "分值", "所属分组ID"])
    ws.append(["哪些是传输层协议？", "TCP\nUDP\nIP\nHTTP", "BA", "", 2, "网络", 2, ""])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    result = parse_workbook(buf)
    row = result.rows[0]
    assert row.valid is True
    assert row.answer == "AB"


def test_parse_judgement_row():
    wb = Workbook()
    ws = wb.active
    ws.title = "判断题"
    ws.append(["题干", "答案", "解析", "难度", "知识点标签", "分值", "所属分组ID"])
    ws.append(["HTTP 默认端口是80", "正确", "", 1, "网络", 1, ""])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    result = parse_workbook(buf)
    assert result.total == 1
    row = result.rows[0]
    assert row.type == "判断题"
    assert row.valid is True
    assert row.answer == "正确"


def test_parse_judgement_accepts_boolean_and_uppercase_cells():
    """Excel 布尔单元格与文本 "TRUE"/"False" 都必须能解析。

    用户输入 TRUE/FALSE 时 Excel 会存成布尔值（openpyxl 读出 Python bool），而原实现
    直接与含小写 "true"/"false" 的字面量集合比较，这两类写法都被判为非法 —— 与单选题
    分支的 `.upper()` 归一化口径不一致。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "判断题"
    ws.append(["题干", "答案", "解析", "难度", "知识点标签", "分值", "所属分组ID"])
    ws.append(["布尔 TRUE", True, "", 1, "网络", 1, ""])
    ws.append(["布尔 FALSE", False, "", 1, "网络", 1, ""])
    ws.append(["文本 TRUE", "TRUE", "", 1, "网络", 1, ""])
    ws.append(["文本 False", "False", "", 1, "网络", 1, ""])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    result = parse_workbook(buf)
    assert [r.answer for r in result.rows] == ["正确", "错误", "正确", "错误"]
    assert all(r.valid for r in result.rows), [r.error for r in result.rows]


def test_parse_fill_row():
    wb = Workbook()
    ws = wb.active
    ws.title = "填空题"
    ws.append(["题干", "答案", "解析", "难度", "知识点标签", "分值", "所属分组ID"])
    ws.append(["TCP三次握手用__包", "SYN/同步", "", 2, "网络", 2, ""])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    result = parse_workbook(buf)
    assert result.total == 1
    row = result.rows[0]
    assert row.type == "填空题"
    assert row.valid is True
    assert row.answer == [["SYN", "同步"]]


def test_parse_fill_blank_count_mismatch_is_row_error():
    """答案空位数必须与题干空位数一致，否则会存下永远判错的题。

    判分要求 `len(correct_answer) == len(user_answer)`（grading._grade_fill），而前端按
    题干里连续下划线的数量渲染输入框：答案少写一个空位时该题永远判错，原实现不报错。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "填空题"
    ws.append(["题干", "答案", "解析", "难度", "知识点标签", "分值", "所属分组ID"])
    ws.append(["甲____与乙____分别是什么？", "甲答案|", "", 2, "网络", 2, ""])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    result = parse_workbook(buf)
    row = result.rows[0]
    assert row.valid is False
    assert "空位数" in row.error


def test_parse_drag_duplicate_left_item_is_row_error():
    """拖拽题左项重复必须报行级错误。

    重复左项会让映射覆盖前一对（`mapping[left] = right`），而 left_items/right_items 仍
    逐行追加，于是「映射数 < 左项数」——判分要求两者长度相等，该题实际不可作答。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "拖拽题"
    ws.append(["题干", "题项与正确容器", "解析", "难度", "知识点标签", "分值", "所属分组ID"])
    ws.append(["HTTP 与 HTTPS 的默认端口", "HTTP:80\nHTTP:443", "", 2, "网络", 2, ""])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    result = parse_workbook(buf)
    row = result.rows[0]
    assert row.valid is False
    assert "重复" in row.error


def test_parse_row_group_id_column():
    """「所属分组ID」列必须被解析（原实现读都不读，模板承诺的按行指定分组不生效）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "单选题"
    ws.append(["题干", "选项", "答案", "解析", "难度", "知识点标签", "分值", "所属分组ID"])
    ws.append(["q1", "A.甲\nB.乙", "A", "", 2, "网络", 2, "7"])
    ws.append(["q2", "A.甲\nB.乙", "A", "", 2, "网络", 2, "abc"])
    ws.append(["q3", "A.甲\nB.乙", "A", "", 2, "网络", 2, ""])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    result = parse_workbook(buf)
    first, second, third = result.rows
    assert first.group_id == 7
    assert second.valid is False and "所属分组ID" in second.error
    assert third.group_id is None  # 留空 = 沿用上传时选择的分组


def test_parse_options_supports_pipe_separator():
    """单行用 | 分隔（无 A. 前缀）时按 | 拆分：原实现把判断放在「结果非空」之后，分支不可达。"""
    from app.utils.excel import _parse_options

    assert _parse_options("甲|乙") == ["甲", "乙"]
    assert _parse_options("A.甲\nB.乙") == ["甲", "乙"]

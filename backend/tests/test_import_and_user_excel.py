"""题库导入与学生/用户 Excel 解析的端到端测试。

`import_service`（覆盖率 20%）与 `user_excel`（21%）此前几乎完全未被执行：
它们承载「Excel 解析 → 预览暂存 → 确认落库」两步流程与 confirm_token 的 IDOR 防护，
是数据入口上最需要回归的路径。
"""

from __future__ import annotations

from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import Workbook

from app.core.errors import DomainError
from app.database import db_session, init_db
from app.models.group import Group
from app.models.question import Question
from app.models.user import User
from app.services import import_service, user_service
from app.utils import user_excel
from app.utils.excel import HEADERS, SHEET_ORDER


def _question_workbook(rows_by_sheet: dict[str, list[list]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in SHEET_ORDER:
        ws = wb.create_sheet(sheet)
        ws.append(HEADERS[sheet])
        for row in rows_by_sheet.get(sheet, []):
            ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _user_workbook(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "用户"
    ws.append(user_excel.HEADERS)
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _admin(db, email: str = "importer@example.com") -> User:
    u = User(email=email, password_hash="x", name="管理员", role="super_admin", status="active", email_verified=True)
    db.add(u)
    db.flush()
    return u


# ---------- 题库导入 ----------
def test_question_import_preview_then_confirm_creates_bank_and_questions():
    init_db()
    with db_session() as db:
        admin = _admin(db)
        content = _question_workbook(
            {
                "单选题": [["HTTP 默认端口?", "A.21\nB.80", "B", "80", 1, "网络", 2, ""]],
                "多选题": [["关系型数据库?", "A.MySQL\nB.Redis", "A", "", 2, "数据库", 3, ""]],
                "判断题": [["HTTP 无状态。", "正确", "", 1, "网络", 2, ""]],
                "填空题": [["TCP 二次握手标志位 ____。", "SYN/同步", "", 2, "网络", 2, ""]],
                "简答题": [["简述 HTTPS。", "HTTP+TLS", "", 3, "安全", 5, ""]],
                "拖拽题": [["协议端口匹配", "HTTP:80\nSSH:22", "", 2, "网络", 3, ""]],
            }
        )
        db.commit()

        preview = import_service.preview(
            db, content, group_id=None, bank_id=None, bank_name="导入库", user_id=admin.id, scope=None
        )
        assert preview["valid_count"] == 6
        assert preview["total"] == 6
        assert preview["truncated"] is False
        token = preview["confirm_token"]
        assert token

        result = import_service.do_import(db, token, user_id=admin.id, scope=None)
        assert result.success == 6
        assert result.failed == 0

        from app.models.question import QuestionBank

        bank = db.query(QuestionBank).filter(QuestionBank.name == "导入库").one()
        assert db.query(Question).filter(Question.bank_id == bank.id).count() == 6
        # 标签自动登记
        from app.models.question import QuestionTag

        assert {t.name for t in db.query(QuestionTag).all()} >= {"网络", "数据库", "安全"}


def test_question_import_reports_invalid_rows_and_existing_bank():
    init_db()
    with db_session() as db:
        admin = _admin(db)
        from app.models.question import QuestionBank

        bank = QuestionBank(name="既有库", group_id=None, practice_enabled=True)
        db.add(bank)
        db.flush()
        content = _question_workbook(
            {
                "单选题": [
                    ["题干正常", "A.a\nB.b", "A", "", 1, "", 2, ""],
                    ["", "A.a", "A", "", 1, "", 2, ""],  # 题干为空 → invalid
                    ["答案越界", "A.a\nB.b", "Z", "", 1, "", 2, ""],  # 答案超出选项范围
                ]
            }
        )
        db.commit()

        preview = import_service.preview(
            db, content, group_id=None, bank_id=bank.id, bank_name="", user_id=admin.id, scope=None
        )
        assert preview["valid_count"] == 1
        assert len(preview["errors"]) == 2

        result = import_service.do_import(db, preview["confirm_token"], user_id=admin.id, scope=None)
        assert result.success == 1
        assert db.query(Question).filter(Question.bank_id == bank.id).count() == 1


def test_question_import_scope_and_owner_and_token_guards():
    init_db()
    with db_session() as db:
        admin = _admin(db)
        own = Group(name="本部门", type="部门")
        other = Group(name="其他部门", type="部门")
        db.add_all([own, other])
        db.flush()
        db.commit()

        with pytest.raises((DomainError, HTTPException)) as forbidden:
            import_service.preview(
                db,
                b"PK\x03\x04",
                group_id=other.id,
                bank_id=None,
                bank_name="",
                user_id=admin.id,
                scope={own.id},
            )
        assert forbidden.value.status_code == 403

        content = _question_workbook({"单选题": [["题", "A.a", "A", "", 1, "", 2, ""]]})
        preview = import_service.preview(
            db, content, group_id=None, bank_id=None, bank_name="库", user_id=admin.id, scope=None
        )
        # 他人 token → 403，且不消费原持有者的预览
        with pytest.raises((DomainError, HTTPException)) as idor:
            import_service.do_import(db, preview["confirm_token"], user_id=admin.id + 999, scope=None)
        assert idor.value.status_code == 403

        assert import_service.do_import(db, preview["confirm_token"], user_id=admin.id, scope=None).success == 1

        # token 已单次消费
        with pytest.raises((DomainError, HTTPException)) as consumed:
            import_service.do_import(db, preview["confirm_token"], user_id=admin.id, scope=None)
        assert consumed.value.status_code == 400

        with pytest.raises((DomainError, HTTPException)) as unknown:
            import_service.do_import(db, "not-a-real-token", user_id=admin.id, scope=None)
        assert unknown.value.status_code == 400


def test_question_import_rejects_non_xlsx_magic_bytes():
    init_db()
    with db_session() as db:
        admin = _admin(db)
        db.commit()
        with pytest.raises(ValueError):
            import_service.preview(
                db, b"not a zip at all", group_id=None, bank_id=None, bank_name="", user_id=admin.id, scope=None
            )


# ---------- 用户 Excel ----------
def test_user_template_preview_and_consume_strips_plaintext_password():
    init_db()
    with db_session() as db:
        admin = _admin(db)
        db.commit()

        content = user_excel.build_template().getvalue()
        preview = user_excel.preview(content, user_id=admin.id)

        # 模板首行示例口令为空 → invalid；次行为部门管理员
        assert preview["valid_count"] == 1
        assert all("password" not in row for row in preview["rows"])

        rows = user_excel.consume_preview(preview["confirm_token"], admin.id)
        assert len(rows) == 1
        assert rows[0]["password"] == ""  # 明文不留在缓存里
        assert rows[0]["password_hash"]

        created = user_service.import_users(db, admin.id, rows, scope=None, actor_role="super_admin")
        assert created["success"] == 1
        assert created["failed"] == 0


@pytest.mark.parametrize(
    ("row", "expected_error"),
    [
        (["", "无名", "普通用户", "Abc12345", "正常", ""], "邮箱为空"),
        (["bad-email", "无名", "普通用户", "Abc12345", "正常", ""], "邮箱格式不正确"),
        (["ok@example.com", "无名", "普通用户", "123", "正常", ""], "初始密码至少 6 位"),
        (["ok@example.com", "无名", "普通用户", "", "正常", ""], "初始密码不能为空"),
        (["ok@example.com", "无名", "不存在的角色", "Abc12345", "正常", ""], "角色非法"),
        (["ok@example.com", "无名", "普通用户", "Abc12345", "不存在的状态", ""], "状态非法"),
    ],
)
def test_user_excel_row_validation_errors(row, expected_error):
    init_db()
    with db_session() as db:
        admin = _admin(db)
        db.commit()

        content = _user_workbook([row])
        preview = user_excel.preview(content, user_id=admin.id)

        assert preview["valid_count"] == 0
        assert len(preview["errors"]) == 1
        assert expected_error in preview["errors"][0]["error"]


def test_user_excel_normalizes_role_status_and_group_ids():
    init_db()
    with db_session() as db:
        admin = _admin(db)
        db.commit()

        content = _user_workbook([["a@example.com", "甲", "部门管理员", "Abc12345", "待审批", "3，5, 999999"]])
        preview = user_excel.preview(content, user_id=admin.id)
        rows = user_excel.consume_preview(preview["confirm_token"], admin.id)

        assert rows[0]["role"] == "dept_admin"
        assert rows[0]["status"] == "pending"
        assert rows[0]["group_ids"] == [3, 5, 999999]  # 解析保留原始 id，存在性由服务层校验
        assert rows[0]["email"] == "a@example.com"

        # 非超管不得导入管理员账号（垂直越权防护）
        created = user_service.import_users(db, admin.id, rows, scope=None, actor_role="user")
        assert created["success"] == 0
        assert created["failed"] == 1


def test_user_excel_preview_guards_unknown_token_and_owner():
    init_db()
    with db_session() as db:
        admin = _admin(db)
        db.commit()

        with pytest.raises((DomainError, HTTPException)) as unknown:
            user_excel.consume_preview("missing", admin.id)
        assert unknown.value.status_code == 400

        content = _user_workbook([["b@example.com", "乙", "普通用户", "Abc12345", "正常", ""]])
        preview = user_excel.preview(content, user_id=admin.id)
        with pytest.raises((DomainError, HTTPException)) as idor:
            user_excel.consume_preview(preview["confirm_token"], admin.id + 999)
        assert idor.value.status_code == 403


def test_user_excel_rejects_invalid_archive():
    init_db()
    with db_session() as db:
        admin = _admin(db)
        db.commit()
        with pytest.raises(ValueError):
            user_excel.preview(b"definitely-not-xlsx", user_id=admin.id)
